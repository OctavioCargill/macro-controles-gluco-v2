"""Pruebas de Netting con datos inventados (no se usan datos reales)."""

from datetime import date, datetime

import openpyxl
import pytest

from netting import cargas
from netting.calendario import dias_a_pagar
from netting.motor import (
    Avisos, Estado, armar_comisiones, armar_propuestas, clasificar_beneficiario, numerar,
    preparar_077, preparar_historico,
)
from netting.salida import asignar_grupos, resumen_propuestas
from netting.servicio import ejecutar_netting, generar_salida

FERIADOS = {date(2026, 11, 23)}


# ---------------------------------------------------------------- días a pagar

@pytest.mark.parametrize("hoy, pago, esperado", [
    (date(2026, 10, 26), datetime(2026, 10, 26), 0),      # hoy
    (date(2026, 10, 26), datetime(2026, 10, 20), 0),      # vencida
    (date(2026, 10, 26), None, 0),                        # sin fecha
    (date(2026, 10, 26), datetime(2026, 10, 27), 1),      # mañana
    (date(2026, 10, 26), datetime(2026, 10, 28), 2),
    (date(2026, 10, 30), datetime(2026, 11, 2), 1),       # viernes → lunes
    (date(2026, 10, 30), datetime(2026, 10, 31), 1),      # viernes → sábado cuenta hasta el lunes
    (date(2026, 11, 20), datetime(2026, 11, 24), 1),      # viernes → martes con lunes feriado
    (date(2026, 10, 26), datetime(2026, 10, 26, 15, 0), 1),  # hoy con hora: la macro cuenta 1
])
def test_dias_a_pagar(hoy, pago, esperado):
    assert dias_a_pagar(pago, hoy, FERIADOS) == esperado


# ---------------------------------------------------------------- tipo de beneficiario

def _h(**kw):
    base = dict(corredor="", vendedor="V", cobrador="V", cesion="", usuario_anulacion="")
    base.update(kw)
    return base


@pytest.mark.parametrize("fila, esperado", [
    (_h(corredor="C", cobrador="C"), "CORREDOR"),
    (_h(corredor="C"), "CORREDOR-DIRECTO"),
    (_h(), "VENDEDOR"),
    (_h(corredor="C", cesion="X"), "CORREDOR-CESION"),
    (_h(cesion="X"), "VENDEDOR-CESION"),
    (_h(usuario_anulacion="CARGILL", corredor="C", cobrador="C"), "CARGILL"),
    (_h(cobrador="OTRO"), ""),
])
def test_clasificar(fila, esperado):
    assert clasificar_beneficiario(fila) == esperado


# ---------------------------------------------------------------- numeración y grupos

def _p(ab, tipo="CORREDOR", tipo_doc="L#", corredor="C", vendedor="V", benef=None):
    return dict(ab_benef=ab, benef=benef or f"B{ab}", tipo_benef=tipo, tipo_doc=tipo_doc,
                corredor=corredor, vendedor=vendedor, n_propuesta="", total=1.0, tipo_pagar="PAGAR")


def test_numerar_24_agrupa_por_ab_y_separa_directos_y_comisiones():
    filas = [
        _p(1), _p(1), _p(1, tipo_doc="C0"), _p(1, tipo_doc="C0"),     # liquidaciones y comisiones del mismo AB
        _p(2, tipo="CORREDOR-DIRECTO", corredor="C", vendedor="V1"),
        _p(2, tipo="CORREDOR-DIRECTO", corredor="C", vendedor="V1"),  # mismo corredor y vendedor: misma propuesta
        _p(2, tipo="CORREDOR-DIRECTO", corredor="C", vendedor="V2"),  # cambia vendedor: otra
        _p(3, tipo="VENDEDOR-CESION"), _p(3, tipo="VENDEDOR-CESION"),  # cesión de vendedor: una por documento
    ]
    estado = Estado([], [], 100, set(), Avisos())
    numerar(filas, estado, "24")
    assert [f["n_propuesta"] for f in filas] == [101, 101, 102, 102, 103, 103, 104, 105, 106]
    assert estado.ultimo_numero == 106


def test_grupos_redondeo_y_tope():
    filas = [dict(n_propuesta=n) for n in (1, 1, 2, 3, 4, 5)]
    assert asignar_grupos(filas, 2, "24") == 5      # 5/2 = 2,5 → 2 (redondeo al par)
    assert [f["grupo"] for f in filas] == [1, 1, 1, 2, 2, 2]  # tope en el último analista
    filas48 = [dict(n_propuesta=n) for n in (1, 2, 3, 4, 5)]
    asignar_grupos(filas48, 2, "48")
    assert [f["grupo"] for f in filas48] == [1, 1, 2, 2, 3]   # 48 hs no tiene tope (como la macro)


# ---------------------------------------------------------------- corrida completa con archivos inventados

HOY = date(2026, 10, 28)  # miércoles
D0, D1, D2, D5 = datetime(2026, 10, 28), datetime(2026, 10, 29), datetime(2026, 10, 30), datetime(2026, 11, 4)

ENCABEZADO_HIST = [
    "Documento", "Tipo Doc", "Descripcion de la liquidacion", "Contrato", "Numero documento fiscal",
    "Usuario anulacion", "Corredor", "Nombre Corredor", "Vendedor", "Nombre Vendedor",
    "AB Beneficiario del Pago", "Nombre Beneficiario del Pago", "AB Beneficiario de Cesion",
    "Nombre Beneficiario de Cesion", "AB Beneficiario Contrato", "Nombre Beneficiario Contrato",
    "Total a Pagar", "Fecha Generacion Liquidacion", "Fecha Pago", "Condicion de Inscripcion",
    "Usuario Alta", "Anticipo/Barter", "ID CORTO", "ID LARGO", "Pago Contra", "Condicion de Pago",
    "Dias A pagar",
]


def _hist(doc, desc, contrato, ab_corr, corr, ab_vend, vend, ab_cob, cob, total, pago, ab_ces="", ces="", aa=""):
    return [doc, "L#" if "Comisi" not in desc else "C0", desc, contrato, "", "", ab_corr, corr, ab_vend, vend,
            ab_cob, cob, ab_ces, ces, 0, "", total, D0, pago, "", "", "", "", f"ID{doc}", "", "", aa]


FILAS_HIST = [
    # Corredor AGRO SUR (AB 500) cobra dos liquidaciones a 24 hs de dos vendedores distintos
    _hist(1001, "Liquidación Final", "K1", 500, "AGRO SUR SA", 900, "CAMPO UNO", 500, "AGRO SUR SA", 1000.0, D1),
    _hist(1002, "Liquidación Final", "K2", 500, "AGRO SUR SA", 901, "CAMPO DOS", 500, "AGRO SUR SA", 2000.0, D0),
    # Vendedor directo LA ESTANCIA (AB 902) cobra a 24 hs
    _hist(1003, "Liquidación Parcial", "K3", "0", "", 902, "LA ESTANCIA", 902, "LA ESTANCIA", 5000.0, D1),
    # Corredor directo: corredor PAMPA (501) pero cobra el vendedor CAMPO TRES (903), 48 hs
    _hist(1004, "Liquidación Final", "K4", 501, "PAMPA CEREALES", 903, "CAMPO TRES", 903, "CAMPO TRES", 700.0, D2),
    # Vendedor con cesión a 24 hs
    _hist(1005, "Liquidación Final", "K5", "0", "", 904, "CAMPO CUATRO", 904, "CAMPO CUATRO", 300.0, D1, 950, "BANCO X"),
    # Comisión del corredor AGRO SUR a 24 hs
    _hist(1006, "Comisiónes", "K1", 500, "AGRO SUR SA", 900, "CAMPO UNO", 500, "AGRO SUR SA", 50.0, D0),
    # Vence en más de 48 hs: queda afuera
    _hist(1007, "Liquidación Final", "K6", 500, "AGRO SUR SA", 900, "CAMPO UNO", 500, "AGRO SUR SA", 999.0, D5),
    # Marcado X en la macro: se descarta
    _hist(1008, "Liquidación Final", "K1", 500, "AGRO SUR SA", 900, "CAMPO UNO", 500, "AGRO SUR SA", 1.0, D0, aa="X"),
]

ENCABEZADO_077 = [
    "Codigo BP Comercial", "BP Comercial", "Cod Cliente", "CLIENTE", "CUIT", "COD TIPO DOC", "TIPO_DOC",
    "NRO DOC", "EMISOR", "NUMERO FACTURA", "NUMERO LEGAL", "DETALLE", "PRODUCTO", "DESC_PRODUCTO", "COSECHA",
    "FECHA_DOC", "VENCIMIENTO_DOC", "DIAS_ATRASO", "VL_LIMITE_CREDITO", "Importe Doc",
    "Saldo pend aplicar Netting", "Intereses", "Intereses COMPENSATORIO", "Reajuste TC", "FECHA_REPORTE",
    "Factura Asociada", "Tipo Factura Asociada", "Contrato", "Corredor", "CUIT Corredor", "AB Corredor",
    "Compania",
]


def _f077(cliente_ab, cliente, tipo, nro, saldo, contrato, corredor="", fecha=datetime(2026, 9, 1)):
    return ["0151195000", "Glucovil", cliente_ab, cliente, "", tipo, "", str(nro), "", "", "", f"Factura {nro}",
            "", "", "", fecha, None, 0, None, saldo, saldo, 0, None, 0, None, "", "", contrato, corredor, "", "", "51511"]


FILAS_077 = [
    _f077(900, "CAMPO UNO", "RI", 7001, 100.0, "K1", "AGRO SUR SA"),       # del corredor, contrato en el 090
    _f077(901, "CAMPO DOS", "AH", 7002, 40.0, "K2", "AGRO SUR SA"),        # del corredor
    _f077(900, "CAMPO UNO", "RI", 7003, 10.0, "K9", "AGRO SUR SA"),        # contrato que no está en el 090: se borra
    _f077(902, "LA ESTANCIA", "RM", 7004, 60.0, "K3"),                     # del vendedor directo
    _f077(902, "LA ESTANCIA", "ZZ", 7005, 60.0, "K3"),                     # tipo fuera de alcance
    _f077(903, "CAMPO TRES", "RI", 7006, 30.0, "K4", "PAMPA CEREALES"),     # corredor directo: no se descuenta
]

ENCABEZADO_090 = ["Contrato"] + [f"c{i}" for i in range(2, 14)] + ["Nombre Vendedor", "c15", "c16", "c17", "Nombre Corredor"]


def _f090(contrato, vendedor, corredor):
    fila = [""] * len(ENCABEZADO_090)
    fila[0], fila[13], fila[17] = contrato, vendedor, corredor
    return fila


FILAS_090 = [
    _f090("K1", "CAMPO UNO", "AGRO SUR SA"),
    _f090("K2", "CAMPO DOS", "AGRO SUR SA"),
    _f090("K3", "LA ESTANCIA", None),
    _f090("K4", "CAMPO TRES", "PAMPA CEREALES"),
    _f090("K5", "CAMPO CUATRO", None),
]

ENCABEZADO_VENTAS = ["FECHA", "AB CLIENTE", "CLIENTE", "CORREDOR", "AB CORREDOR", "DOC;INTERNO", "DOCUMENTO",
                     "ESTADO", "IMPORTE A COMPENSAR", "IMPORTE COMPENSADO", "DIFERNECIA", "DIAS DE ESPERA"]
FILAS_VENTAS = [
    [45000, "35100001", "COMPRADOR SA", "AGRO SUR SA", "500", "QJ 1802", 1, "PENDIENTE", 80.0, 0, 80.0, None],
    [45000, "35100002", "OTRO SA", "AGRO SUR SA", "500", "ESPERAR", 1, "PENDIENTE", 15.0, 0, 15.0, None],
]


def _xlsx(ruta, encabezado, filas, hoja="Hoja1", filas_titulo=0):
    libro = openpyxl.Workbook()
    hoja_activa = libro.active
    hoja_activa.title = hoja
    for _ in range(filas_titulo):
        hoja_activa.append(["Reporte SAP"])
    hoja_activa.append(encabezado)
    for fila in filas:
        hoja_activa.append(fila)
    libro.save(ruta)
    return str(ruta)


@pytest.fixture
def archivos(tmp_path):
    return {
        "historico": _xlsx(tmp_path / "Historico.xlsx", ENCABEZADO_HIST, FILAS_HIST),
        "077": _xlsx(tmp_path / "077.xlsx", ENCABEZADO_077, FILAS_077 + [["Data Refresh Date: hoy"]], filas_titulo=2),
        "090": _xlsx(tmp_path / "090.xlsx", ENCABEZADO_090, FILAS_090, filas_titulo=2),
        "ventas": _xlsx(tmp_path / "ventas.xlsx", ENCABEZADO_VENTAS, FILAS_VENTAS, hoja="Compensas Pendientes"),
    }


CONFIG = {"feriados": ["2026-01-01", "2026-11-23"], "corredores_solo_24": [], "rutas": {}}


def test_corrida_completa(archivos, tmp_path, monkeypatch):
    resultado = ejecutar_netting(archivos, HOY, 1000, CONFIG)
    p24, p48 = resultado.propuestas_24, resultado.propuestas_48

    def docs(filas, tipo):
        return sorted(f["doc"] for f in filas if f["tipo"] == tipo)

    # 24 hs: liquidaciones 1001, 1002, 1003, 1005 y la comisión 1006; no entran 1004 (48), 1007 (>48), 1008 (X)
    assert docs(p24, "PAGAR") == [1001, 1002, 1003, 1005, 1006]
    # Descuentos: 7001 y 7002 del corredor (con contrato en el 090), la venta 1802, y 7004 del vendedor directo
    assert docs(p24, "COBRAR") == [1802, 7001, 7002, 7004]
    cobrar = {f["doc"]: f for f in p24 if f["tipo"] == "COBRAR"}
    assert cobrar[7001]["total"] == -100.0          # el 077 se descuenta con signo negativo
    assert cobrar[1802]["total"] == 80.0            # la venta toma la diferencia tal cual (como la macro)
    assert cobrar[7001]["benef"] == "AGRO SUR SA"

    # Propuestas: AGRO SUR (liquidaciones + descuentos), LA ESTANCIA, cesión, comisión AGRO SUR aparte
    resumen = resultado.resumen("24")
    por_benef = {(r["benef"], r["filas"][0]["tipo_doc"]): r for r in resumen}
    agro = por_benef[("AGRO SUR SA", "L#")]
    assert agro["a_pagar"] == 3000.0 and agro["a_descontar"] == -140.0 + 80.0
    assert por_benef[("LA ESTANCIA", "L#")]["neto"] == 5000.0 - 60.0
    assert por_benef[("BANCO X", "L#")]["documentos"] == 1
    assert [r["n_propuesta"] for r in resumen] == list(range(1001, 1001 + len(resumen)))

    # 48 hs: el corredor directo 1004, sin descuentos
    assert docs(p48, "PAGAR") == [1004] and docs(p48, "COBRAR") == []
    assert resultado.numero_final == 1000 + len(resumen) + 1

    # Salida
    monkeypatch.setattr("netting.servicio.guardar_config", lambda c: None)
    ruta = generar_salida(resultado, 2, tmp_path, dict(CONFIG))
    libro = openpyxl.load_workbook(ruta)
    assert libro.sheetnames == ["Propuestas_24", "Propuestas_48", "Control"]
    hoja = libro["Propuestas_24"]
    assert [c.value for c in hoja[1]][14] == "N_PROPUESTA"
    assert hoja.max_row == len(p24) + 1


def test_corrida_sin_compensas_de_ventas(archivos):
    sin_ventas = dict(archivos, ventas="")
    resultado = ejecutar_netting(sin_ventas, HOY, 1000, CONFIG)
    cobrar = sorted(f["doc"] for f in resultado.propuestas_24 if f["tipo"] == "COBRAR")
    assert cobrar == [7001, 7002, 7004]  # sin la venta 1802
    assert any("SIN compensas de ventas" in a for a in resultado.avisos.revisar)


def test_errores_de_carga_claros(archivos, tmp_path):
    with pytest.raises(cargas.ErrorDeCarga, match="faltan columnas obligatorias"):
        cargas.cargar_historico(_xlsx(tmp_path / "malo.xlsx", ["Documento", "Tipo Doc", "Total a Pagar", "Fecha Pago"], [[1, "L#", 1, D0]]))
    with pytest.raises(cargas.ErrorDeCarga, match="Compensas Pendientes"):
        cargas.cargar_ventas(archivos["077"])
