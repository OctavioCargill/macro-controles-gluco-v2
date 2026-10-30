"""Lectura de los tres archivos de entrada de Netting.

Se leen las celdas tal cual (sin que pandas cambie tipos) para poder replicar las comparaciones
de la macro. Se aceptan .xlsb (pyxlsb), .xlsx/.xlsm (openpyxl), .xls y .csv.

    Histórico   el que genera CONTROLAR o el de la macro de control vieja; se ubican las
                columnas por nombre, así sirve cualquiera de los dos formatos.
    077         reporte de documentos pendientes de cobro (fila de encabezado detectada sola).
    090         reporte de contratos de compra.
"""

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path


class ErrorDeCarga(Exception):
    """Error con un mensaje pensado para mostrarle al usuario."""


# ---------------------------------------------------------------- utilidades de celdas

def vacio(valor) -> bool:
    """Equivalente a `celda = Empty` en VBA: vacío, texto vacío o el número 0."""
    if valor is None:
        return True
    if isinstance(valor, str):
        return valor == ""
    if isinstance(valor, bool):
        return False
    if isinstance(valor, (int, float)):
        return valor == 0
    return False


def normalizar(valor):
    """Unifica lo que devuelven pyxlsb y openpyxl: None→"", 35109575.0→35109575."""
    if valor is None:
        return ""
    if isinstance(valor, float):
        if valor != valor:  # NaN
            return ""
        if valor.is_integer():
            return int(valor)
    return valor


_EPOCA_EXCEL = datetime(1899, 12, 30)


def a_fecha(valor):
    """Convierte un valor de celda a datetime como lo haría Excel/VBA (dd/mm/aaaa)."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if valor <= 0:
            return None
        return _EPOCA_EXCEL + timedelta(days=float(valor))
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def a_numero(valor):
    """Como CDbl: número o None si no se puede convertir."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        return float(valor)
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    if re.fullmatch(r"-?[\d.]+,\d+", texto):  # 1.234,56
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", "")
    try:
        return float(texto)
    except ValueError:
        return None


def a_numero_si_se_puede(valor):
    """Como Texto en columnas: '1802' pasa a 1802; el resto queda igual."""
    if isinstance(valor, str):
        texto = valor.strip()
        if re.fullmatch(r"-?\d+", texto):
            return int(texto)
        if re.fullmatch(r"-?\d+\.\d+", texto):
            return normalizar(float(texto))
    return valor


# ---------------------------------------------------------------- lectura de hojas

def _leer_xlsb(ruta: Path, hoja=None):
    try:
        from pyxlsb import open_workbook
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise ErrorDeCarga("Falta la librería pyxlsb para leer archivos .xlsb.") from exc
    with open_workbook(str(ruta)) as libro:
        nombres = list(libro.sheets)
        if hoja is not None and hoja not in nombres:
            raise ErrorDeCarga(
                f"El archivo {ruta.name} no tiene la hoja «{hoja}». Hojas disponibles: {', '.join(nombres)}"
            )
        with libro.get_sheet(hoja if hoja is not None else 1) as sheet:
            return [[celda.v for celda in fila] for fila in sheet.rows(sparse=False)]


def _leer_openpyxl(ruta: Path, hoja=None):
    import openpyxl

    libro = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    try:
        if hoja is not None:
            if hoja not in libro.sheetnames:
                raise ErrorDeCarga(
                    f"El archivo {ruta.name} no tiene la hoja «{hoja}». Hojas disponibles: {', '.join(libro.sheetnames)}"
                )
            sheet = libro[hoja]
        else:
            sheet = libro.active or libro.worksheets[0]
        return [list(fila) for fila in sheet.iter_rows(values_only=True)]
    finally:
        libro.close()


def _leer_csv(ruta: Path):
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = ruta.read_text(encoding=codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    filas = []
    for fila in csv.reader(texto.splitlines(), dialecto):
        filas.append([a_numero_si_se_puede(c) if c != "" else None for c in fila])
    return filas


def _leer_xls(ruta: Path, hoja=None):
    import pandas as pd

    try:
        df = pd.read_excel(ruta, sheet_name=hoja if hoja is not None else 0, header=None, dtype=object)
    except ValueError as exc:
        raise ErrorDeCarga(f"No se pudo leer la hoja «{hoja}» de {ruta.name}.") from exc
    return [[None if (v is None or v != v) else v for v in fila] for fila in df.itertuples(index=False)]


def leer_hoja(ruta, hoja=None):
    ruta = Path(ruta)
    if not ruta.exists():
        raise ErrorDeCarga(f"No se encontró el archivo: {ruta}")
    sufijo = ruta.suffix.lower()
    try:
        if sufijo == ".xlsb":
            filas = _leer_xlsb(ruta, hoja)
        elif sufijo in (".xlsx", ".xlsm"):
            filas = _leer_openpyxl(ruta, hoja)
        elif sufijo == ".csv":
            filas = _leer_csv(ruta)
        elif sufijo == ".xls":
            filas = _leer_xls(ruta, hoja)
        else:
            raise ErrorDeCarga(f"Formato no soportado: {ruta.name}")
    except ErrorDeCarga:
        raise
    except PermissionError as exc:
        raise ErrorDeCarga(f"No se pudo abrir {ruta.name}: está abierto por otro programa o no hay permisos.") from exc
    except Exception as exc:
        raise ErrorDeCarga(f"No se pudo leer {ruta.name}: {exc}") from exc
    return [[normalizar(v) for v in fila] for fila in filas]


def _clave(texto):
    return re.sub(r"\s+", " ", str(texto or "")).strip().casefold()


def _buscar_encabezado(filas, obligatorias, nombre_archivo, max_filas=20):
    claves = [_clave(o) for o in obligatorias]
    for indice, fila in enumerate(filas[:max_filas]):
        valores = {_clave(v) for v in fila}
        if all(c in valores for c in claves):
            return indice
    raise ErrorDeCarga(
        f"{nombre_archivo}: no se encontró la fila de encabezados (se buscaron las columnas "
        f"{', '.join(obligatorias)} en las primeras {max_filas} filas). ¿Es el archivo correcto?"
    )


def _mapa_columnas(encabezado, columnas, nombre_archivo):
    """columnas: {campo: [nombres posibles]} → {campo: índice}. Falla si falta alguna obligatoria."""
    posiciones = {}
    for indice, valor in enumerate(encabezado):
        clave = _clave(valor)
        if clave and clave not in posiciones:
            posiciones[clave] = indice
    mapa, faltantes = {}, []
    for campo, (alternativas, obligatoria) in columnas.items():
        for nombre in alternativas:
            if _clave(nombre) in posiciones:
                mapa[campo] = posiciones[_clave(nombre)]
                break
        else:
            if obligatoria:
                faltantes.append(alternativas[0])
    if faltantes:
        raise ErrorDeCarga(f"{nombre_archivo}: faltan columnas obligatorias: {', '.join(faltantes)}.")
    return mapa


def _valor(fila, indice):
    if indice is None or indice >= len(fila):
        return ""
    return fila[indice]


@dataclass
class Lectura:
    filas: list
    avisos: list = field(default_factory=list)
    ruta: str = ""


# ---------------------------------------------------------------- Histórico

COLUMNAS_HISTORICO = {
    # campo: ([nombres posibles], obligatoria)   (letra en la hoja Historico de la macro)
    "doc": (["Documento", "Numero Doc"], True),                          # A
    "tipo_doc": (["Tipo Doc"], True),                                     # B
    "desc": (["Descripcion de la liquidacion"], True),                    # C
    "contrato": (["Contrato"], True),                                     # D
    "usuario_anulacion": (["Usuario anulacion"], False),                  # F (la macro la pisa)
    "ab_corredor": (["Corredor"], True),                                  # G
    "corredor": (["Nombre Corredor"], True),                              # H
    "ab_vendedor": (["Vendedor"], True),                                  # I
    "vendedor": (["Nombre Vendedor"], True),                              # J
    "ab_cobrador": (["AB Beneficiario del Pago"], True),                  # K
    "cobrador": (["Nombre Beneficiario del Pago"], True),                 # L
    "ab_cesion": (["AB Beneficiario de Cesion"], True),                   # M
    "cesion": (["Nombre Beneficiario de Cesion"], True),                  # N
    "total": (["Total a Pagar"], True),                                   # Q
    "fecha_pago": (["Fecha Pago"], True),                                 # S
    "dias_origen": (["Dias A pagar"], False),                             # AA (marca PENDIENTE / X)
    "id_largo": (["ID LARGO", "id_largo"], False),
}


def cargar_historico(ruta) -> Lectura:
    filas = leer_hoja(ruta)
    nombre = Path(ruta).name
    inicio = _buscar_encabezado(filas, ["Tipo Doc", "Total a Pagar", "Fecha Pago"], f"Histórico ({nombre})")
    mapa = _mapa_columnas(filas[inicio], COLUMNAS_HISTORICO, f"Histórico ({nombre})")
    registros, avisos = [], []
    datos = filas[inicio + 1:]
    for numero, fila in enumerate(datos, start=inicio + 2):
        registro = {campo: _valor(fila, indice) for campo, indice in mapa.items()}
        if vacio(registro["doc"]):
            restantes = sum(1 for f in datos[numero - inicio - 1:] if not vacio(_valor(f, mapa["doc"])))
            if restantes:
                avisos.append(
                    f"Histórico: la fila {numero} no tiene número de documento; como la macro, se ignoran "
                    f"las {restantes} filas con datos que siguen."
                )
            break
        for campo in COLUMNAS_HISTORICO:
            registro.setdefault(campo, "")
        registro["fecha_pago"] = a_fecha(registro["fecha_pago"]) if registro["fecha_pago"] != "" else None
        registro["fila_origen"] = numero
        registros.append(registro)
    if not registros:
        raise ErrorDeCarga(f"Histórico ({nombre}): no tiene documentos.")
    return Lectura(registros, avisos, str(ruta))


# ---------------------------------------------------------------- 077

COLUMNAS_077 = {
    "cod_bp": (["Codigo BP Comercial"], False),            # A
    "cod_cliente": (["Cod Cliente"], True),                # C
    "cliente": (["CLIENTE"], True),                        # D
    "tipo": (["COD TIPO DOC"], True),                      # F
    "nro": (["NRO DOC"], True),                            # H
    "detalle": (["DETALLE"], True),                        # L
    "fecha": (["FECHA_DOC"], True),                        # P
    "saldo": (["Saldo pend aplicar Netting"], True),       # U
    "contrato": (["Contrato"], True),                      # AB
    "corredor": (["Corredor"], True),                      # AC (nombre)
    "ab_corredor": (["AB Corredor"], False),               # AE
}


def cargar_077(ruta) -> Lectura:
    filas = leer_hoja(ruta)
    nombre = Path(ruta).name
    inicio = _buscar_encabezado(filas, ["COD TIPO DOC", "NRO DOC"], f"077 ({nombre})")
    mapa = _mapa_columnas(filas[inicio], COLUMNAS_077, f"077 ({nombre})")
    registros, avisos = [], []
    datos = filas[inicio + 1:]
    for posicion, fila in enumerate(datos):
        registro = {campo: _valor(fila, indice) for campo, indice in mapa.items()}
        if vacio(registro["tipo"]):
            restantes = sum(
                1 for f in datos[posicion + 1:] if not vacio(_valor(f, mapa["tipo"]))
            )
            if restantes:
                avisos.append(
                    f"077: la fila {inicio + 2 + posicion} no tiene tipo de documento; se ignoran las "
                    f"{restantes} filas con datos que siguen."
                )
            break
        for campo in COLUMNAS_077:
            registro.setdefault(campo, "")
        registro["fila_origen"] = inicio + 2 + posicion
        registros.append(registro)
    return Lectura(registros, avisos, str(ruta))


# ---------------------------------------------------------------- 090

COLUMNAS_090 = {
    "contrato": (["Contrato"], True),             # A
    "vendedor": (["Nombre Vendedor"], True),      # N
    "corredor": (["Nombre Corredor"], True),      # R
}


def cargar_090(ruta) -> Lectura:
    filas = leer_hoja(ruta)
    nombre = Path(ruta).name
    inicio = _buscar_encabezado(filas, ["Contrato", "Nombre Vendedor", "Nombre Corredor"], f"090 ({nombre})")
    mapa = _mapa_columnas(filas[inicio], COLUMNAS_090, f"090 ({nombre})")
    registros = []
    for numero, fila in enumerate(filas[inicio + 1:], start=inicio + 2):
        registro = {campo: _valor(fila, indice) for campo, indice in mapa.items()}
        if all(vacio(v) for v in registro.values()):
            continue
        registro["fila_origen"] = numero
        registros.append(registro)
    if not registros:
        raise ErrorDeCarga(f"090 ({nombre}): no tiene contratos.")
    return Lectura(registros, [], str(ruta))
